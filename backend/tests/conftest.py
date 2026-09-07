"""
conftest.py
===========
Shared fixtures.

Only the PDF is checked in — a text-bearing PDF is awkward to build without
pulling in a PDF writer the app does not otherwise need. The .docx and .pptx
fixtures are generated at test time from ``python-docx`` and ``python-pptx``,
which are already runtime dependencies, so no binary files are committed for
formats we can create in three lines.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

# Tests are run from backend/ (``pytest``) but also from the repo root by some
# editors; make ``app`` importable either way.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def pdf_path() -> Path:
    """A three-page, text-bearing PDF."""
    return FIXTURES / "sample.pdf"


@pytest.fixture(scope="session")
def docx_path(tmp_path_factory) -> Path:
    """A .docx with headings, paragraphs and a table."""
    import docx

    path = tmp_path_factory.mktemp("docs") / "sample.docx"
    document = docx.Document()
    document.add_heading("Cell Biology", 0)
    document.add_paragraph("Mitochondria generate ATP via oxidative phosphorylation.")
    # Contains a zero-width space, which the cleaner must strip.
    document.add_paragraph("The nucleus stores DNA.​  Ribosomes translate mRNA.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Organelle"
    table.cell(0, 1).text = "Function"
    table.cell(1, 0).text = "Ribosome"
    table.cell(1, 1).text = "Protein synthesis"

    document.save(path)
    return path


@pytest.fixture(scope="session")
def pptx_path(tmp_path_factory) -> Path:
    """A three-slide deck."""
    from pptx import Presentation

    path = tmp_path_factory.mktemp("decks") / "sample.pptx"
    deck = Presentation()
    for index, title in enumerate(["Intro to Genetics", "Mendel's Laws", "Punnett Squares"]):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Key point {index + 1}: alleles segregate independently."
    deck.save(path)
    return path


@pytest.fixture
def serve(monkeypatch):
    """Patch ``download`` to return bytes from a name → content mapping.

    Shared because both the extraction tests and the retrieval tests need to
    drive the pipeline without a network. A name mapped to an exception raises
    it instead, which is how a failing download is simulated.
    """
    from app.services import extraction

    def _serve(mapping):
        def fake_download(url, *, file_name, max_bytes):
            value = mapping[file_name]
            if isinstance(value, Exception):
                raise value
            return value

        monkeypatch.setattr(extraction, "download", fake_download)

    return _serve
