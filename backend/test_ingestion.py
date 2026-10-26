"""Offline format fixtures, including real OLE containers for DOC/PPT readers."""
import io
import struct
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import pymupdf
from PIL import Image
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches
from app.services.rag.ingestion.manager import IngestionManager
from app.services.rag.ingestion.detector import FileTypeDetector
from app.services.rag.ingestion.exceptions import UnsupportedFileTypeError
from app.services.rag.schemas import SupportedFormat
from app.services.rag.chunking.manager import ChunkingManager
from app.services.rag.chunking.models import ChunkingConfig
from app.services.rag.processing.cleaning.cleaner import TextCleaner

TEXT = "Cells are the basic units of life. DNA carries genetic information."


def ole_container(streams: dict[str, bytes]) -> bytes:
    free, end, fat_marker = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    names = sorted(streams)
    assert len(names) <= 3
    sectors = [bytearray(512)]
    entries = []
    for name in names:
        data = streams[name].ljust(max(4096, ((len(streams[name]) + 511) // 512) * 512), b'\0')
        first = len(sectors)
        sectors.extend(bytearray(data[i:i + 512]) for i in range(0, len(data), 512))
        entries.append((name, first, len(data)))
    fat_sector = len(sectors)
    fat = [free] * 128
    fat[0] = end
    for _, first, size in entries:
        count = size // 512
        for sector in range(first, first + count):
            fat[sector] = sector + 1 if sector < first + count - 1 else end
    fat[fat_sector] = fat_marker
    sectors.append(bytearray(struct.pack('<128I', *fat)))
    def directory(name, kind, first, size, right=free, child=free):
        record = bytearray(128)
        encoded = (name + '\0').encode('utf-16le')
        record[:len(encoded)] = encoded
        struct.pack_into('<HBBIII', record, 64, len(encoded), kind, 1, free, right, child)
        struct.pack_into('<IQ', record, 116, first, size)
        return record
    sectors[0][:128] = directory('Root Entry', 5, end, 0, child=1)
    for index, (name, first, size) in enumerate(entries, 1):
        sectors[0][index * 128:(index + 1) * 128] = directory(name, 2, first, size, right=index + 1 if index < len(entries) else free)
    header = bytearray(512)
    header[:8] = bytes.fromhex('D0CF11E0A1B11AE1')
    struct.pack_into('<HHHHH', header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into('<IIIIIIIII', header, 40, 0, 1, 0, 0, 4096, end, 0, end, 0)
    struct.pack_into('<109I', header, 76, fat_sector, *([free] * 108))
    return bytes(header) + b''.join(sectors)


def legacy_doc() -> bytes:
    word = bytearray(4096)
    struct.pack_into('<HH', word, 0, 0xA5EC, 0x00C1)
    struct.pack_into('<II', word, 24, 512, 512 + len(TEXT) * 2)
    struct.pack_into('<I', word, 0x4C, len(TEXT))
    piece = struct.pack('<II', 0, len(TEXT)) + struct.pack('<HIH', 0, 512, 0)
    clx = b'\x02' + struct.pack('<I', len(piece)) + piece
    struct.pack_into('<II', word, 0x1A2, 0, len(clx))
    word[512:512 + len(TEXT) * 2] = TEXT.encode('utf-16le')
    return ole_container({'WordDocument': bytes(word), '0Table': clx})


def legacy_ppt() -> bytes:
    payload = TEXT.encode('utf-16le')
    atom = struct.pack('<HHI', 0, 4000, len(payload)) + payload
    slide = struct.pack('<HHI', 15, 1006, len(atom)) + atom
    return ole_container({'PowerPoint Document': slide})


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.manager = IngestionManager()

    def test_modern_and_legacy_formats_are_identified_from_signatures(self):
        word = WordDocument(); word.add_paragraph(TEXT); docx = io.BytesIO(); word.save(docx)
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]); slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = TEXT
        pptx = io.BytesIO(); deck.save(pptx)
        pdf = pymupdf.open(); pdf.new_page().insert_text((40, 50), TEXT); pdf_bytes = pdf.tobytes(); pdf.close()
        fixtures = [(pdf_bytes, SupportedFormat.PDF), (docx.getvalue(), SupportedFormat.DOCX),
                    (pptx.getvalue(), SupportedFormat.PPTX), (legacy_doc(), SupportedFormat.DOC), (legacy_ppt(), SupportedFormat.PPT)]
        for content, expected in fixtures:
            with self.subTest(format=expected):
                result = self.manager.ingest_bytes(content, 'renamed.wrong', owner_id='owner')
                self.assertEqual(result.file_type, expected.value)
                self.assertIn('genetic information', '\n'.join(d.text for d in result.documents))
                self.assertEqual(result.file_name, 'renamed.wrong')

    def test_image_reader_uses_vision_ocr(self):
        buffer = io.BytesIO(); Image.new('RGB', (50, 50), 'white').save(buffer, format='PNG')
        with patch('app.services.rag.ingestion.loaders.image_loader.read_image', return_value=TEXT) as ocr:
            result = self.manager.ingest_bytes(buffer.getvalue(), 'lecture.png', owner_id='owner')
        self.assertEqual(result.documents[0].text, TEXT)
        self.assertEqual(result.file_type, 'image')
        ocr.assert_called_once()

    def test_scanned_pdf_uses_vision_ocr(self):
        buffer = io.BytesIO(); Image.new('RGB', (200, 100), 'white').save(buffer, format='PNG')
        pdf = pymupdf.open(); page = pdf.new_page(); page.insert_image(page.rect, stream=buffer.getvalue()); content = pdf.tobytes(); pdf.close()
        with patch('app.services.rag.ingestion.loaders.pdf_loader.read_image', return_value=TEXT) as ocr:
            result = self.manager.ingest_bytes(content, 'scan.pdf', owner_id='owner')
        self.assertIn('genetic', result.documents[0].text)
        ocr.assert_called_once()

    def test_renaming_non_pdf_does_not_make_it_a_pdf(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fake.pdf'; path.write_bytes(b'not a PDF')
            with self.assertRaises(UnsupportedFileTypeError):
                FileTypeDetector.detect(path)

    def test_reuploads_have_stable_owner_scoped_source_and_chunk_ids(self):
        first = self.manager.ingest_bytes(legacy_doc(), 'original.doc', 'owner-a')
        second = self.manager.ingest_bytes(legacy_doc(), 'original.doc', 'owner-a')
        other = self.manager.ingest_bytes(legacy_doc(), 'original.doc', 'owner-b')
        self.assertEqual(first.source_id, second.source_id)
        self.assertNotEqual(first.source_id, other.source_id)
        chunker = ChunkingManager(ChunkingConfig(tokenizer_fn=lambda text: len(text.split())))
        a, b = chunker.chunk_documents(first.documents), chunker.chunk_documents(second.documents)
        self.assertEqual([n.node_id for n in a.parent_chunks + a.child_chunks], [n.node_id for n in b.parent_chunks + b.child_chunks])

    def test_line_wrap_and_hyphen_repair_preserve_paragraphs_and_code(self):
        text = 'This is a long prose line about how genetic information is\ncarried by DNA.\n\nAn inter-\nnational example.\n\n```\ncode-\nword\n```'
        result = TextCleaner().clean(text)
        self.assertIn('is carried', result)
        self.assertIn('international', result)
        self.assertIn('code-\nword', result)
        self.assertIn('\n\n', result)

    def test_page_cleanup_preserves_numeric_code_and_body_values(self):
        text = 'Measured value:\n42\nunits\n\n```\n123\n```\n\nPage 4 of 5\n5'
        result = TextCleaner().clean(text)
        self.assertIn('42', result)
        self.assertIn('```\n123\n```', result)
        self.assertNotIn('Page 4 of 5', result)
        self.assertFalse(result.rstrip().endswith('5'))
