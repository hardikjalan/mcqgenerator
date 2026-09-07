"""
pptx_loader.py
==============
Loader for modern PowerPoint files (.pptx — Office Open XML).

Uses ``python-pptx`` directly to iterate over slides, shapes, and tables,
preserving slide titles as Markdown headings (``##``), body text, tables
as Markdown grids, and speaker notes where available and accurately
extractable.  Each slide becomes a separate ``Document``.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from llama_index.core.schema import Document

from app.services.rag.schemas import SupportedFormat, StandardDocumentMetadata
from app.services.rag.ingestion.base import BaseLoader
from app.services.rag.ingestion.registry import loader_for
from app.services.rag.ingestion.exceptions import CorruptedFileError


def _table_to_markdown(table) -> str:
    """Convert a ``python-pptx`` table to a Markdown grid.

    Returns an empty string if the table has no rows.
    """
    rows = list(table.rows)
    if not rows:
        return ""

    lines: list[str] = []
    for r_idx, row in enumerate(rows):
        cells = [cell.text.replace("|", "\\|").strip() for cell in row.cells]
        lines.append("| " + " | ".join(cells) + " |")
        if r_idx == 0:
            lines.append("| " + " | ".join("---" for _ in cells) + " |")

    return "\n".join(lines)


@loader_for(SupportedFormat.PPTX)
class PptxLoader(BaseLoader):
    """Read a .pptx file and return one Document per slide.

    Uses ``python-pptx`` directly to extract slide titles as Markdown
    headings, body text from content placeholders, tables as Markdown
    grids, and speaker notes — preserving structural formatting where
    available.
    """

    def load(self, file_path: Path, metadata: StandardDocumentMetadata) -> list[Document]:
        try:
            prs = Presentation(str(file_path))
        except Exception as exc:
            raise CorruptedFileError(
                file_path.name,
                reason=f"python-pptx could not read this file: {exc}",
            ) from exc

        total_slides = len(prs.slides)
        result: list[Document] = []

        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            parts: list[str] = []

            # Extract slide title as a Markdown heading
            if slide.shapes.title and slide.shapes.title.text.strip():
                parts.append(f"## {slide.shapes.title.text.strip()}")

            # Extract text from all other shapes (body placeholders, text boxes)
            for shape in slide.shapes:
                # Skip the title shape — already handled above
                if shape == slide.shapes.title:
                    continue

                if shape.has_table:
                    md_table = _table_to_markdown(shape.table)
                    if md_table:
                        parts.append(md_table)

                elif shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            parts.append(text)

            # Extract speaker notes
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    parts.append(f"\n**Notes:** {notes_text}")

            text = "\n\n".join(parts)

            doc = Document(text=text)
            slide_meta = metadata.model_copy(
                update={
                    "page_or_slide_num": slide_num,
                    "total_pages_or_slides": total_slides,
                    "custom_metadata": {
                        "slide_number": slide_num,
                    },
                }
            )
            self._attach_metadata(doc, slide_meta)
            result.append(doc)

        return self.clean_and_validate(result, file_name=file_path.name)
